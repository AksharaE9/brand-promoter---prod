import os
import re
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import win32com.client

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=60, bottom=60, left=90, right=90):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(f'''
            <w:tblBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="6" w:space="0" w:color="{color}"/>
                <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{color}"/>
                <w:left w:val="none"/>
                <w:right w:val="none"/>
                <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
                <w:insideV w:val="none"/>
            </w:tblBorders>
        ''')
        tblPr[0].append(borders)

def build_full_perf_docx(md_path, docx_path):
    doc = Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        
        footer = section.footer
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        f_run = f_p.add_run("TalentOS Full Performance & Load Test Audit | Internal Engineering Report")
        f_run.font.name = 'Calibri'
        f_run.font.size = Pt(8.5)
        f_run.font.color.rgb = RGBColor(140, 150, 160)

    with open(md_path, 'r', encoding='utf-8') as f:
        md_text = f.read()

    lines = md_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        
        if not line:
            i += 1
            continue

        # Horizontal Rule
        if re.match(r'^---+', line):
            i += 1
            continue

        # Header 1
        if line.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(line.replace('# ', '').strip())
            r.font.name = 'Calibri'
            r.font.size = Pt(18)
            r.font.bold = True
            r.font.color.rgb = RGBColor(26, 54, 93)
            i += 1
            continue

        # Header 2
        if line.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(line.replace('## ', '').strip())
            r.font.name = 'Calibri'
            r.font.size = Pt(13)
            r.font.bold = True
            r.font.color.rgb = RGBColor(43, 108, 176)
            i += 1
            continue

        # Header 3
        if line.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(line.replace('### ', '').strip())
            r.font.name = 'Calibri'
            r.font.size = Pt(10.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(30, 41, 59)
            i += 1
            continue

        # Blockquote
        if line.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.left_indent = Inches(0.2)
            clean_q = line.replace('> ', '').replace('**', '').strip()
            r = p.add_run(clean_q)
            r.font.name = 'Calibri'
            r.font.size = Pt(9.5)
            r.font.italic = True
            r.font.color.rgb = RGBColor(71, 85, 105)
            i += 1
            continue

        # Code Block
        if line.startswith('```'):
            i += 1
            code_lines = []
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i])
                i += 1
            if i < len(lines):
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.left_indent = Inches(0.15)
            r = p.add_run('\n'.join(code_lines))
            r.font.name = 'Consolas'
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor(30, 41, 59)
            continue

        # Table
        if line.startswith('|'):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            
            # parse table
            rows_data = []
            for tl in table_lines:
                if '---' in tl:
                    continue
                cells = [c.strip() for c in tl.split('|')[1:-1]]
                if cells:
                    rows_data.append(cells)
            
            if rows_data:
                num_cols = len(rows_data[0])
                table = doc.add_table(rows=len(rows_data), cols=num_cols)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.autofit = True
                
                # Header
                for c_idx, cell_text in enumerate(rows_data[0]):
                    cell = table.rows[0].cells[c_idx]
                    set_cell_background(cell, "1A365D")
                    set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(0)
                    clean_text = cell_text.replace('**', '').replace('&rarr;', '->').replace('&times;', 'x')
                    r = p.add_run(clean_text)
                    r.font.name = 'Calibri'
                    r.font.bold = True
                    r.font.size = Pt(8)
                    r.font.color.rgb = RGBColor(255, 255, 255)

                # Body
                for r_idx in range(1, len(rows_data)):
                    bg_color = "FFFFFF" if r_idx % 2 != 0 else "F8FAFC"
                    for c_idx in range(min(num_cols, len(rows_data[r_idx]))):
                        cell = table.rows[r_idx].cells[c_idx]
                        set_cell_background(cell, bg_color)
                        set_cell_margins(cell, top=40, bottom=40, left=80, right=80)
                        p = cell.paragraphs[0]
                        p.paragraph_format.space_after = Pt(0)
                        cell_text = rows_data[r_idx][c_idx].replace('**', '').replace('&rarr;', '->').replace('&times;', 'x')
                        r = p.add_run(cell_text)
                        r.font.name = 'Calibri'
                        r.font.size = Pt(8)
                        if c_idx == 0:
                            r.font.bold = True
                            r.font.color.rgb = RGBColor(30, 41, 59)
                        elif "PASS" in cell_text or "MET" in cell_text or "100%" in cell_text:
                            r.font.bold = True
                            r.font.color.rgb = RGBColor(22, 101, 52)
                        elif "BLOCKED" in cell_text or "WARN" in cell_text:
                            r.font.bold = True
                            r.font.color.rgb = RGBColor(180, 83, 9)
                        else:
                            r.font.color.rgb = RGBColor(71, 85, 105)

                set_table_borders(table, "CBD5E1")
            continue

        # Bullet List
        if line.startswith('- ') or line.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            clean_b = line[2:].replace('**', '').replace('&rarr;', '->').replace('&times;', 'x')
            r = p.add_run(clean_b)
            r.font.name = 'Calibri'
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(51, 65, 85)
            i += 1
            continue

        # Numbered List
        if re.match(r'^\d+\.\s+', line):
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            clean_n = re.sub(r'^\d+\.\s+', '', line).replace('**', '').replace('&rarr;', '->')
            r = p.add_run(clean_n)
            r.font.name = 'Calibri'
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(51, 65, 85)
            i += 1
            continue

        # Normal Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        clean_p = line.replace('**', '').replace('&rarr;', '->').replace('&times;', 'x')
        r = p.add_run(clean_p)
        r.font.name = 'Calibri'
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(51, 65, 85)
        i += 1

    doc.save(docx_path)
    print(f"Generated DOCX at {docx_path}")

def convert_to_pdf_safe(docx_path, pdf_path):
    word = win32com.client.Dispatch('Word.Application')
    word.Visible = False
    try:
        doc = word.Documents.Open(docx_path)
        doc.SaveAs(pdf_path, FileFormat=17) # 17 = wdFormatPDF
        doc.Close()
        print(f"Generated PDF at {pdf_path}")
    except Exception as e:
        print(f"PDF generation note: {e}")
    finally:
        try:
            word.Quit()
        except:
            pass

if __name__ == '__main__':
    base_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), 'reports/perf'))
    md_file = os.path.join(base_dir, 'FINAL-PERFORMANCE-REPORT.md')
    docx_file = os.path.join(base_dir, 'FINAL-PERFORMANCE-REPORT.docx')
    pdf_file = os.path.join(base_dir, 'FINAL-PERFORMANCE-REPORT.pdf')
    
    build_full_perf_docx(md_file, docx_file)
    convert_to_pdf_safe(docx_file, pdf_file)
